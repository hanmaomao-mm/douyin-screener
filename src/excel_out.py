# -*- coding: utf-8 -*-
"""Excel 输出：每条候选/刷量记录实时追加写盘，可随时打开查看。"""
from __future__ import annotations

import os
import threading
import time
from collections import deque
from datetime import datetime
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import CONFIG

HEADERS = [
    ("序号", 6),
    ("视频链接", 34),
    ("达人昵称", 18),
    ("达人主页链接", 34),
    ("看后搜词条", 22),
    ("话题标签", 26),
    ("命中关键词", 16),
    ("判定理由", 52),
    ("播放量", 11),
    ("点赞数", 11),
    ("评论数", 11),
    ("点赞率", 12),
    ("等级标记", 11),
    ("识别时间", 19),
    ("截图存档", 30),
]

_GRADE_FILL = {
    "高匹配": PatternFill("solid", fgColor="FFE3E3"),   # 红系（符合国内涨/高热语义）
    "普通": PatternFill("solid", fgColor="FFF6E0"),
    "刷量号": PatternFill("solid", fgColor="E4F3E4"),   # 绿系（表示剔除）
    "不达标": PatternFill("solid", fgColor="F0F0F0"),
}


class ExcelWriter:
    """追加式 Excel 写入器。

    去重口径：在「最近 N 条」的滑动窗口内，同一个达人只记一次
    （N = config.yaml 的 safety.dedup_memory）。加载已有文件时会按同样窗口预热，
    所以重启工作流不会把同一个达人重复写进去。
    """

    def __init__(self, path: str, sheet_name: str = "候选达人", freeze_header: bool = True,
                 dedup_memory: int = 400):
        self.path = path
        self.sheet_name = sheet_name
        self.freeze_header = freeze_header
        self.dedup_memory = max(1, int(dedup_memory))
        self._lock = threading.RLock()
        self._seq = 0
        self._recent: deque[str] = deque(maxlen=self.dedup_memory)
        self._seen: set[str] = set()
        self._pending_warn = ""
        os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
        existed = os.path.isfile(self.path)
        self._wb, self._ws = self._open_or_create()
        if not existed:
            # 首次运行就先把带表头的空表落盘，方便你随时打开查看
            self.save()

    # ---------------------------------------------------------- 初始化
    def _open_or_create(self):
        if os.path.isfile(self.path):
            try:
                wb = load_workbook(self.path)
                ws = wb[self.sheet_name] if self.sheet_name in wb.sheetnames else wb.create_sheet(self.sheet_name)
                self._seq = max(0, ws.max_row - 1)
                # 用最后 dedup_memory 行的达人昵称预热去重窗口
                start = max(2, ws.max_row - self.dedup_memory + 1)
                for row in ws.iter_rows(min_row=start, values_only=True):
                    if row and len(row) > 2 and row[2]:
                        self._remember(str(row[2]))
                return wb, ws
            except Exception:
                pass
        wb = Workbook()
        ws = wb.active
        ws.title = self.sheet_name
        self._write_header(ws)
        return wb, ws

    def _remember(self, nickname: str) -> None:
        key = nickname.strip().casefold()
        if not key:
            return
        if len(self._recent) == self._recent.maxlen and self._recent:
            old = self._recent[0]
            if old not in list(self._recent)[1:]:
                self._seen.discard(old)
        self._recent.append(key)
        self._seen.add(key)

    def _write_header(self, ws) -> None:
        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for idx, (name, width) in enumerate(HEADERS, start=1):
            c = ws.cell(row=1, column=idx, value=name)
            c.font = Font(bold=True, color="FFFFFF", size=11)
            c.fill = PatternFill("solid", fgColor="2F5597")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.row_dimensions[1].height = 24
        if self.freeze_header:
            ws.freeze_panes = "A2"

    @staticmethod
    def _key(nickname: str) -> str:
        return nickname.strip().casefold()

    # ---------------------------------------------------------- 写入
    def exists(self, nickname: str) -> bool:
        with self._lock:
            return self._key(nickname) in self._seen

    def append(self, row: Dict[str, Any], nickname: str, frame_ts: str = "") -> bool:
        """追加一行。返回 True 表示写入成功（False = 该达人在去重窗口内已记录过）。"""
        with self._lock:
            key = self._key(nickname)
            if nickname and key in self._seen:
                return False
            self._seq += 1
            r = self._seq + 1
            ws = self._ws

            values = [
                self._seq,
                row.get("video_link", ""),
                row.get("nickname", ""),
                row.get("profile_link", ""),
                row.get("search_terms", ""),
                row.get("hashtags", ""),
                row.get("hit_keywords", ""),
                row.get("reason", ""),
                row.get("play", ""),
                row.get("like", ""),
                row.get("comment", ""),
                row.get("like_rate", ""),
                row.get("grade", ""),
                row.get("ts", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                row.get("frame_path", ""),
            ]
            thin = Side(style="thin", color="E8E8E8")
            for ci, v in enumerate(values, start=1):
                c = ws.cell(row=r, column=ci, value=v)
                c.alignment = Alignment(vertical="center", wrap_text=(ci == 8))
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            # 超链接
            for ci, url in ((2, row.get("video_link", "")), (4, row.get("profile_link", ""))):
                if isinstance(url, str) and url.startswith("http"):
                    cell = ws.cell(row=r, column=ci)
                    cell.hyperlink = url
                    cell.font = Font(color="0563C1", underline="single")
            grade = str(row.get("grade", ""))
            fill = _GRADE_FILL.get(grade)
            if fill:
                ws.cell(row=r, column=13).fill = fill
                ws.cell(row=r, column=13).font = Font(bold=True)

            if nickname:
                self._remember(nickname)
            self.save()
            return True

    def save(self) -> None:
        """写盘。若 Excel 正被打开会导致失败，这里做重试并给出提示。"""
        for attempt in range(3):
            try:
                self._wb.save(self.path)
                self._pending_warn = ""
                return
            except PermissionError:
                self._pending_warn = "Excel 文件被占用（可能正开着），请关闭后自动重试"
                time.sleep(1.2)
            except Exception as exc:
                self._pending_warn = f"保存 Excel 失败：{exc}"
                return
        print(f"[警告] {self._pending_warn}")

    @property
    def count(self) -> int:
        return self._seq

    def last_warning(self) -> str:
        return self._pending_warn
